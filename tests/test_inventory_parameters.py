import csv
import json
import math
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from preservation_helpers import Profile, assert_preserved, export_snapshots, import_artifact, synthetic_snapshot

from brightpath import BackgroundContext, BiosphereProfile, FormatProfile, InventoryContext, TechnosphereProfile
from brightpath.analysis import analyze_inventory

PROFILE = Profile("uvek", "2025")
SUPPLIER = {
    "name": "Electricity, low voltage, at grid",
    "reference product": "Electricity, low voltage, at grid",
    "location": "CH",
    "unit": "kilowatt hour",
}


@pytest.mark.parametrize("source_format", ["brightway", "openlca", "simapro"])
@pytest.mark.parametrize("target_format", ["brightway", "openlca", "simapro"])
def test_parameters_survive_all_format_pairs(source_format, target_format):
    source = synthetic_snapshot(SUPPLIER)
    before = deepcopy(source)
    artifact = export_snapshots([source], PROFILE, source_format)
    restored = import_artifact(artifact, PROFILE)
    assert_preserved(source, restored[0])
    exported = export_snapshots(restored, PROFILE, target_format)
    assert_preserved(source, import_artifact(exported, PROFILE)[0])
    assert source == before


def test_openlca_calculated_parameters_and_native_uncertainty():
    snapshot = synthetic_snapshot(SUPPLIER)
    snapshot["parameters"][0].update({"uncertainty type": 2, "loc": math.log(2), "scale": 0.2})
    artifact = export_snapshots([snapshot], PROFILE, "openlca")
    with ZipFile(BytesIO(artifact.content)) as archive:
        entities = {name: json.loads(archive.read(name)) for name in archive.namelist() if name.endswith(".json")}
    process = next(value for name, value in entities.items() if name.startswith("processes/"))
    parameters = {parameter["name"]: parameter for parameter in process["parameters"]}
    assert parameters["dose"]["isInputParameter"] is False
    assert parameters["dose"]["formula"] == "shared * efficiency"
    assert not parameters["dose"].get("uncertainty")
    assert parameters["efficiency"]["isInputParameter"] is True
    uncertainty = parameters["efficiency"]["uncertainty"]
    assert uncertainty["geomMean"] == pytest.approx(2)
    assert uncertainty["geomSd"] == pytest.approx(math.exp(0.2))
    assert_preserved(snapshot, import_artifact(artifact, PROFILE)[0])


def test_excel_formulas_are_literal_text_not_executable_cells():
    artifact = export_snapshots([synthetic_snapshot(SUPPLIER)], PROFILE, "brightway")
    workbook = load_workbook(BytesIO(artifact.content), data_only=False)
    cells = [cell for sheet in workbook for row in sheet for cell in row]
    assert any(cell.value == "shared * efficiency" for cell in cells)
    assert not any(cell.data_type == "f" for cell in cells)


@pytest.mark.parametrize(
    "delimiter,suffix,format_id", [(",", ".csv", "brightway_csv"), ("\t", ".tsv", "brightway_tsv")]
)
def test_delimited_analysis_retains_shared_parameter_scopes(tmp_path, delimiter, suffix, format_id):
    source = synthetic_snapshot(SUPPLIER)
    artifact = export_snapshots([source], PROFILE, "brightway")
    workbook = load_workbook(BytesIO(artifact.content), data_only=False)
    path = tmp_path / f"inventory{suffix}"
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, delimiter=delimiter)
        for row in workbook.active.iter_rows(values_only=True):
            writer.writerow(["" if value is None else value for value in row])
    context = InventoryContext(
        FormatProfile(format_id),
        BackgroundContext(TechnosphereProfile("uvek", "2025", "cutoff"), BiosphereProfile("ecoinvent", "3.10")),
    )
    result = analyze_inventory(path=path, source_context=context)
    assert not result.has_errors, result.file_issues
    assert result.database_parameters == source["database parameters"]
    assert result.project_parameters == source["project parameters"]
