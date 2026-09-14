"""Apply scoped follow-up decisions and audit the provenance of matching targets."""

from __future__ import annotations

import argparse
import copy
import csv
import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

if __package__:
    from .review_uvek_mapping import FIELDS, identity, unit_key
else:
    from review_uvek_mapping import FIELDS, identity, unit_key

CURATION_FILES = (
    "mapping_decisions.json",
    "manual_overrides.json",
    "reviewed_from_unmapped_3.10.json",
    "reconciled_curated_3.10.csv",
)
FIT_CONFIDENCE = {"exact": 1.0, "good proxy": 0.7, "rough proxy": 0.4}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def group_id(row):
    key = (
        row["reference product"],
        json.loads(row["previous_target"])["name"],
        json.loads(row["candidate_target"])["name"],
        row["outcome"],
    )
    return hashlib.sha256(json.dumps(key).encode()).hexdigest()[:12]


def scope_hash(rows):
    scoped = [
        (list(identity(row)), json.loads(row["previous_target"]), json.loads(row["candidate_target"])) for row in rows
    ]
    return hashlib.sha256(json.dumps(sorted(scoped), sort_keys=True).encode()).hexdigest()


class BranchLineage:
    """Replay branch lookup precedence without treating its labels as validation."""

    def __init__(self, rows_by_file):
        self.lookups = []
        for filenames in ((CURATION_FILES[0],), CURATION_FILES[1:3], (CURATION_FILES[3],)):
            full, inherited = {}, {}
            for filename in filenames:
                for decision in rows_by_file[filename]:
                    source = identity(decision["source"])
                    entry = filename, decision
                    full[source[:3]] = entry
                    inherited.setdefault(source[:2], entry)
            self.lookups.append((full, inherited))

    @classmethod
    def load(cls, directory):
        rows_by_file = {}
        for filename in CURATION_FILES:
            path = directory / filename
            if path.suffix == ".json":
                rows_by_file[filename] = json.loads(path.read_text())["mappings"]
            else:
                with path.open(newline="") as handle:
                    rows_by_file[filename] = [
                        {
                            side: dict(
                                zip(
                                    FIELDS,
                                    (row[f"{side}_{field}"] for field in ("name", "refprod", "location", "unit")),
                                )
                            )
                            for side in ("source", "target")
                        }
                        for row in csv.DictReader(handle)
                    ]
        return cls(rows_by_file)

    def trace(self, row):
        result = {
            "lineage_kind": "algorithmic",
            "lineage_scope": "not_applicable",
            "lineage_file": "",
            "lineage_source": "",
            "lineage_comment": "",
            "lineage_fit": "",
            "lineage_confidence": "",
            "lineage_flags": "",
        }
        if row["candidate_method"] not in ("review", "curated"):
            return result
        source = identity(row)
        for full, inherited in self.lookups:
            entry = full.get(source[:3]) or inherited.get(source[:2])
            if entry is None:
                continue
            filename, decision = entry
            original = identity(decision["source"])
            scope = "exact_source" if source == original else "inherited_location"
            flags = []
            if source[3] != original[3]:
                scope = "inherited_unit"
                flags.append("source_unit_not_checked_by_branch_lookup")
            comment = decision.get("comment", "") or ""
            prefix = comment.strip().lower().split(":", 1)[0]
            kind = "script_" + prefix if prefix in ("rule", "reuse", "fix") else "recorded_decision"
            if filename.endswith(".csv"):
                kind = "auto_reconciled_curated"
            fit = (decision.get("fit") or "").strip().lower()
            confidence = FIT_CONFIDENCE.get(fit, decision.get("confidence", 1.0))
            if fit and fit not in FIT_CONFIDENCE:
                flags.append("fit_label_not_recognized_by_branch_generator")
            if scope != "exact_source":
                flags.append("choice_inherited_without_exact_source_review")
                if float(row["candidate_confidence"]) > confidence:
                    flags.append("inherited_confidence_exceeds_origin")
            if identity(decision["target"]) != identity(json.loads(row["candidate_target"])):
                flags.append("workbook_target_differs_from_lineage")
                kind = "unresolved_lineage"
            expected_method = "curated" if filename.endswith(".csv") else "review"
            if row["candidate_method"] != expected_method:
                flags.append("workbook_method_differs_from_lineage")
            return {
                "lineage_kind": kind,
                "lineage_scope": scope,
                "lineage_file": filename,
                "lineage_source": json.dumps(decision["source"], sort_keys=True),
                "lineage_comment": comment,
                "lineage_fit": fit,
                "lineage_confidence": confidence,
                "lineage_flags": "; ".join(flags),
            }
        return result | {"lineage_kind": "unresolved_lineage", "lineage_flags": "no_recorded_origin_found"}


def identity_risks(row):
    """Return screening flags, not an exhaustive equivalence assessment."""
    source = row["reference product"].lower()
    target = json.loads(row["candidate_target"])
    target_name = target["name"].lower()
    flags = []
    if target["location"] != row["location"]:
        flags.append("geographic_proxy")
    if target_name.startswith("xx"):
        flags.append("deprecated_target")
    if unit_key(target["unit"]) != unit_key(row["unit"]):
        flags.append("unit_mismatch")
    if "1000 units" in target_name and "1000" not in source:
        flags.append("bundled_reference_amount_needs_conversion")
    if source.startswith("electricity") and "allocation heat" in target_name:
        flags.append("electricity_allocation_to_heat_unverified")
    if source.startswith("operation, computer"):
        if any(state in source for state in ("off mode", "standby", "sleep mode")):
            flags.append("off_standby_operation_not_validated")
        if "desktop" in source and "smartphone" in target_name:
            flags.append("desktop_to_smartphone_mismatch")
    if any(word in source for word in ("sludge", "bottom ash", "fly ash", "residues")):
        flags.append("waste_composition_and_mass_basis_unverified")
    if target_name.startswith("chemicals "):
        flags.append("generic_chemical_proxy")
    return "; ".join(flags)


def reviewed_decision(row, adjudication, provenance):
    cap = adjudication["confidence_cap"]
    if isinstance(cap, bool) or not isinstance(cap, (int, float)) or not 0 <= cap <= 0.9:
        raise ValueError("Accepted follow-up requires a finite evidence-specific confidence cap")
    if provenance["lineage_confidence"] != "" and provenance["lineage_kind"] != "unresolved_lineage":
        cap = min(cap, provenance["lineage_confidence"])
    limitations = [
        "Identity-level comparison only; inventory composition and LCIA equivalence are not verified.",
        adjudication["reason"],
    ]
    risks = identity_risks(row)
    if risks:
        limitations.append("Screening flags: " + risks)
    if provenance["lineage_scope"].startswith("inherited"):
        limitations.append("Branch choice inherited from another source identity; not independent exact-source review.")
    if provenance["lineage_kind"].startswith("script_"):
        limitations.append(
            "Branch review label includes a scripted rule, reuse or fix; not independent human validation."
        )
    if row["unit"] != json.loads(row["candidate_target"])["unit"]:
        limitations.append("Equivalent passenger-distance unit spellings require an explicit conversion factor of 1.0.")
    return {
        "source": {field: row[field] for field in FIELDS},
        "previous_target": json.loads(row["previous_target"]),
        "target": json.loads(row["candidate_target"]),
        "previous_confidence": float(row["previous_confidence"]),
        "confidence": min(cap, float(row["candidate_confidence"])),
        "mapping_method": "clic_followup_review",
        "review_outcome": adjudication["outcome"],
        "reason": adjudication["reason"],
        "limitations": limitations,
        "branch_method": row["candidate_method"],
        "branch_reasoning": row["candidate_reasoning"],
        "branch_lineage": provenance,
    }


def followup(resource, overrides, comparison, ledger, lineage, targets):
    """Return reviewed overrides and audit rows, without mutating inputs."""
    if ledger["source_revision"] != overrides["source_revision"]:
        raise ValueError("Follow-up revision differs from the previous review")
    low = {identity(rule["source"]): rule for rule in resource["replace"] if rule["confidence"] < 0.5}
    pending, same = defaultdict(list), []
    seen = set()
    for row in comparison:
        source = identity(row)
        if source in seen:
            raise ValueError("Duplicate source in comparison")
        seen.add(source)
        if source not in low:
            continue
        if row["outcome"] == "same_target":
            same.append(row)
        elif row["outcome"].startswith("review_") and row["outcome"] not in ("review_drop", "review_unmapped"):
            pending[group_id(row)].append(row)
        else:
            continue
        if low[source]["target"] != json.loads(row["previous_target"]) or low[source]["confidence"] != float(
            row["previous_confidence"]
        ):
            raise ValueError(f"Stale comparison for {source}")
    if seen != {identity(rule["source"]) for rule in resource["replace"]}:
        raise ValueError("Comparison must cover every current source identity")
    adjudications = {group["id"]: group for group in ledger["groups"]}
    if len(adjudications) != len(ledger["groups"]) or set(adjudications) != set(pending):
        raise ValueError("Follow-up ledger does not cover exactly the pending groups")
    selected_sources = {identity(decision["source"]) for decision in overrides["decisions"]}
    decisions, audit = [], []
    for identifier, rows in pending.items():
        adjudication = adjudications[identifier]
        if len(rows) != adjudication["entries"] or scope_hash(rows) != adjudication["scope_sha256"]:
            raise ValueError(f"Stale follow-up group {identifier}")
        outcome = adjudication["outcome"]
        if not outcome.startswith(("accept_", "retain_", "reject_", "needs_")):
            raise ValueError(f"Invalid follow-up outcome {outcome}")
        for row in rows:
            provenance = lineage.trace(row)
            confidence = float(row["previous_confidence"])
            if outcome.startswith("accept_"):
                source = identity(row)
                target = identity(json.loads(row["candidate_target"]))
                if source in selected_sources or low[source]["mapping_method"] != "heuristic_similarity":
                    raise ValueError("Follow-up would overwrite a prior reviewed choice")
                if target not in targets or unit_key(row["unit"]) != unit_key(target[3]):
                    raise ValueError("Invalid target or missing unit conversion")
                if "1000 units" in target[0].lower():
                    raise ValueError("Bundled target requires reference-amount review")
                decision = reviewed_decision(row, adjudication, provenance)
                decisions.append(decision)
                confidence = decision["confidence"]
                selected_sources.add(source)
            audit.append(
                row
                | provenance
                | {
                    "followup_outcome": outcome,
                    "followup_reason": adjudication["reason"],
                    "final_confidence": confidence,
                    "screening_flags": identity_risks(row),
                }
            )
    agreement = []
    for row in same:
        if json.loads(row["candidate_target"]) != json.loads(row["previous_target"]):
            raise ValueError("Same-target audit contains different targets")
        agreement.append(
            row
            | lineage.trace(row)
            | {
                "screening_flags": identity_risks(row),
                "followup_outcome": "same_target_score_unchanged",
                "followup_reason": "Agreement and provenance are not independent inventory validation or score calibration.",
                "final_confidence": float(row["previous_confidence"]),
            }
        )
    result = copy.deepcopy(overrides)
    result["decisions"].extend(decisions)
    result["review_version"] = 3
    result["followup_policy"] = (
        "Scoped identity-level adjudication of low-confidence alternatives; same-target provenance audit without score inflation. See docs/uvek_followup_decisions.json."
    )
    counts = dict(sorted(Counter(row["followup_outcome"] for row in audit).items()))
    remaining_low = len(low) - sum(decision["confidence"] >= 0.5 for decision in decisions)
    summary = {
        "source_revision": overrides["source_revision"],
        "review_version": 3,
        "total_source_entries": len(resource["replace"]),
        "low_confidence_before": len(low),
        "low_confidence_after": remaining_low,
        "alternatives_reviewed": len(audit),
        "additional_targets_selected": len(decisions),
        "additional_selected_still_below_0.5": sum(decision["confidence"] < 0.5 for decision in decisions),
        "cumulative_targets_selected": len(result["decisions"]),
        "outcomes": counts,
        "same_target_entries": len(agreement),
        "same_target_score_changes": 0,
        "alternatives_retained": sum(count for outcome, count in counts.items() if outcome.startswith("retain_")),
        "alternatives_rejected": sum(count for outcome, count in counts.items() if outcome.startswith("reject_")),
        "alternatives_need_inventory_evidence": sum(
            count for outcome, count in counts.items() if outcome.startswith("needs_")
        ),
    }
    for label, rows in (
        ("same_target", agreement),
        ("same_target_review_label", [row for row in agreement if row["candidate_method"] == "review"]),
    ):
        for column in ("lineage_kind", "lineage_scope"):
            summary[label + "_" + column] = dict(sorted(Counter(row[column] for row in rows).items()))
        for column in ("lineage_flags", "screening_flags"):
            summary[label + "_" + column] = dict(
                sorted(Counter(flag for row in rows for flag in row[column].split("; ") if flag).items())
            )
    return result, audit, agreement, summary


def write_reports(directory, audit, agreement, summary):
    directory.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook(write_only=True)
    overview = workbook.create_sheet("Summary")
    overview.append(["Metric", "Value"])
    for key, value in summary.items():
        overview.append([key, json.dumps(value, sort_keys=True) if isinstance(value, dict) else value])
    for title, filename, rows in (
        ("Alternatives", "uvek_followup_review.csv", audit),
        ("Same target provenance", "uvek_same_target_provenance.csv", agreement),
    ):
        rows = sorted(rows, key=lambda row: (float(row["previous_confidence"]), identity(row)))
        if not rows:
            continue
        headers = list(rows[0])
        with (directory / filename).open("w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers)
            writer.writeheader()
            writer.writerows(rows)
        sheet = workbook.create_sheet(title)
        sheet.freeze_panes = "A2"
        sheet.append(headers)
        for column, header in enumerate(headers, 1):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = (
                55 if any(word in header for word in ("name", "target", "reason", "comment", "source")) else 22
            )
        for row in rows:
            cells = []
            for header in headers:
                value = row[header]
                cell = openpyxl.cell.WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cells.append(cell)
            sheet.append(cells)
        sheet.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}{len(rows) + 1}"
    workbook.save(directory / "uvek_followup_review.xlsx")
    (directory / "uvek_followup_summary.json").write_text(json.dumps(summary, indent=2, sort_keys=True) + "\n")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ("resource", "overrides", "comparison", "curation_directory", "output_directory"):
        parser.add_argument(name, type=Path)
    root = Path(__file__).resolve().parents[1]
    parser.add_argument("--ledger", type=Path, default=root / "docs/uvek_followup_decisions.json")
    args = parser.parse_args()
    with args.comparison.open(newline="") as handle:
        comparison = list(csv.DictReader(handle))
    catalog = json.loads((root / "brightpath/data/export/reference_catalogs/uvek__2025__cutoff.json").read_text())
    targets = {(row["name"], row["reference_product"], row["location"], row["unit"]) for row in catalog["technosphere"]}
    result, audit, agreement, summary = followup(
        json.loads(args.resource.read_text()),
        json.loads(args.overrides.read_text()),
        comparison,
        json.loads(args.ledger.read_text()),
        BranchLineage.load(args.curation_directory),
        targets,
    )
    result["followup_inputs"] = {
        name: digest(getattr(args, name)) for name in ("resource", "overrides", "comparison", "ledger")
    }
    result["followup_inputs"]["curation"] = {
        filename: digest(args.curation_directory / filename) for filename in CURATION_FILES
    }
    summary["input_hashes"] = result["followup_inputs"]
    write_reports(args.output_directory, audit, agreement, summary)
    (args.output_directory / "uvek_reviewed_overrides.json").write_text(
        json.dumps(result, indent=2, sort_keys=True) + "\n"
    )
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
