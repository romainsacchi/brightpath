"""Conservatively select CLIC workbook improvements and report every comparison."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

import openpyxl

if __package__:
    from .uvek_pairwise import decide
else:
    from uvek_pairwise import decide

FIELDS = ("name", "reference product", "location", "unit")
METHODS = {"review", "curated", "structural_refprod", "structural_v2v3"}
MATERIAL_CORRECTIONS = {
    ("lithium", "Lithium chloride, at plant"),
    ("polycarbonate", "Polycarbonate disk"),
    ("bronze", "Contour, bronze"),
    ("brass", "Contour, brass"),
    ("manganese", "Manganese concentrate, at beneficiation"),
    ("molybdenum", "Molybdenum concentrate, main product"),
}


def identity(record):
    return tuple(record[field] for field in FIELDS)


def product_key(text):
    text = text.lower().strip()
    text = re.sub(r", at .*$", "", text)
    text = re.sub(r", production$", "", text)
    return tuple(sorted(re.findall(r"[a-z0-9]+", text)))


def unit_key(text):
    return re.sub(r"[\s-]+", " ", text.lower().strip())


def disposal(text):
    return text.lower().startswith(("disposal", "treatment", "market for waste", "market for used"))


def select_initial(current, candidate):
    if candidate.get("method") is None:
        return "review_unmapped"
    if candidate["method"] == "dropped":
        return "review_drop"
    source = current["source"]
    target = candidate["target"]
    previous = current["target"]
    if identity(previous) == identity(target):
        return "same_target"
    if current["mapping_method"] != "heuristic_similarity":
        return "retain_existing_curated"
    if source["reference product"].lower().startswith(("electricity", "heat", "steam")):
        return "review_energy_technology"
    if unit_key(source["unit"]) != unit_key(target["unit"]):
        return "review_unit_mismatch"
    if disposal(source["name"]) != disposal(target["name"]):
        return "review_process_mismatch"
    if re.match(r"^xx", target["name"], re.I):
        return "review_deprecated_target"
    if candidate["method"] not in METHODS or candidate["confidence"] < 0.7:
        return "review_candidate_method"
    source_product = product_key(source["reference product"])
    if not source_product or product_key(target["reference product"]) != source_product:
        return "review_product_mismatch"
    if product_key(previous["reference product"]) == source_product and disposal(previous["name"]) == disposal(
        source["name"]
    ):
        return "retain_existing_product_match"
    if (
        set(source_product) <= set(product_key(previous["reference product"]))
        and disposal(previous["name"]) == disposal(source["name"])
        and (source["reference product"], previous["name"]) not in MATERIAL_CORRECTIONS
    ):
        return "review_existing_specific_product"
    return "accept_product_match"


def evaluate(current, candidate):
    if candidate.get("method") is None:
        return "review_unmapped", "Branch has no committed mapping for this identity.", None
    if candidate["method"] == "dropped":
        return (
            "review_drop",
            "Branch proposes no equivalent; keep a visible review flag rather than deleting the exchange.",
            None,
        )
    if identity(current["target"]) == identity(candidate["target"]):
        return "same_target", "Both mappings select the identical target.", None
    if current["mapping_method"] != "heuristic_similarity":
        return "retain_existing_curated", "Preserve the existing curated or previously reviewed decision.", None
    if unit_key(current["source"]["unit"]) != unit_key(candidate["target"]["unit"]):
        return "review_unit_mismatch", "Candidate requires a unit conversion not established by this review.", None
    same_product = all(
        current["target"][field] == candidate["target"][field] for field in ("name", "reference product", "unit")
    )
    if same_product and candidate["target"]["location"] == current["source"]["location"]:
        return (
            "accept_geography",
            "Identical target product and process; branch matches the source geography whereas the previous target does not.",
            0.9,
        )
    if select_initial(current, candidate) == "accept_product_match":
        return (
            "accept_product_match",
            "Previously inspected exact-product/material-form correction from the initial review.",
            0.9,
        )
    return decide(current, candidate)


def select(current, candidate):
    return evaluate(current, candidate)[0]


def review(workbook_path, resource_path, output_directory, revision):
    resource = json.loads(resource_path.read_text())
    catalog_path = (
        Path(__file__).resolve().parents[1] / "brightpath/data/export/reference_catalogs/uvek__2025__cutoff.json"
    )
    catalog = json.loads(catalog_path.read_text())
    catalog_targets = {
        (row["name"], row["reference_product"], row["location"], row["unit"]) for row in catalog["technosphere"]
    }
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    candidates = {}
    for sheet in ("mapping", "unmapped"):
        rows = iter(workbook[sheet].values)
        headers = next(rows)
        for values in rows:
            row = dict(zip(headers, values))
            source = dict(zip(FIELDS, (row["source_" + field] for field in ("name", "refprod", "location", "unit"))))
            if identity(source) in candidates:
                raise ValueError(f"Duplicate source: {source}")
            row["target"] = dict(
                zip(FIELDS, (row.get("target_" + field) for field in ("name", "refprod", "location", "unit")))
            )
            candidates[identity(source)] = row
    workbook.close()
    decisions = []
    audit = []
    for current in resource["replace"]:
        candidate = candidates.get(identity(current["source"]))
        outcome, evidence, confidence_cap = (
            evaluate(current, candidate)
            if candidate
            else ("outside_3.10_catalog", "No source identity in the branch's 3.10 workbook.", None)
        )
        if (
            candidate
            and candidate.get("method") not in (None, "dropped")
            and identity(candidate["target"]) not in catalog_targets
        ):
            outcome, evidence, confidence_cap = (
                "reject_catalog_target",
                "Branch target is not present in the exact packaged UVEK catalog.",
                None,
            )
        audit.append(
            {
                **{field: current["source"][field] for field in FIELDS},
                "outcome": outcome,
                "previous_confidence": current["confidence"],
                "candidate_confidence": candidate.get("confidence", "") if candidate else "",
                "candidate_deprecated": bool(
                    candidate and str(candidate["target"].get("name") or "").lower().startswith("xx")
                ),
                "candidate_geographic_proxy": bool(
                    candidate
                    and candidate["target"].get("location")
                    and candidate["target"]["location"] != current["source"]["location"]
                ),
                "evidence": evidence,
                "previous_target": json.dumps(current["target"], sort_keys=True),
                "candidate_target": json.dumps(candidate["target"], sort_keys=True) if candidate else "",
                "candidate_method": candidate.get("method", "") if candidate else "",
                "candidate_reasoning": candidate.get("reasoning", "") if candidate else "",
            }
        )
        if outcome.startswith("accept_"):
            decisions.append(
                {
                    "source": current["source"],
                    "previous_target": current["target"],
                    "target": candidate["target"],
                    "confidence": min(confidence_cap, candidate["confidence"]),
                    "mapping_method": (
                        "clic_product_checked" if outcome == "accept_product_match" else "clic_pairwise_review"
                    ),
                    "previous_confidence": current["confidence"],
                    "review_outcome": outcome,
                    "limitations": [
                        "Identity-level comparison only; inventory composition and LCIA equivalence are not verified."
                    ]
                    + (
                        ["Deprecated target remains present in the exact UVEK catalog."]
                        if str(candidate["target"]["name"]).lower().startswith("xx")
                        else []
                    )
                    + (
                        ["Geographic proxy."]
                        if candidate["target"]["location"] != current["source"]["location"]
                        else []
                    ),
                    "branch_method": candidate["method"],
                    "branch_reasoning": candidate.get("reasoning", ""),
                    "reason": evidence,
                }
            )
    output_directory.mkdir(parents=True, exist_ok=True)
    payload = {
        "schema_version": 1,
        "source_repository": "https://github.com/romainsacchi/clic",
        "source_revision": revision,
        "source_workbook": workbook_path.name,
        "source_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        "baseline_resource_sha256": hashlib.sha256(resource_path.read_bytes()).hexdigest(),
        "review_version": 2,
        "scope": "Exact source identities present in the CLIC ecoinvent 3.10 cutoff workbook; shared identities inherit this choice across existing resource profiles.",
        "policy": "Pairwise product, energy carrier/fuel/technology, waste material/route, nutrient-basis and co-product checks. Generic chemical proxies can replace weak wrong-product matches; confidence values are not calibrated across inputs. No drops or unresolved candidates are imported.",
        "decisions": decisions,
    }
    (output_directory / "uvek_reviewed_overrides.json").write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n")
    with (output_directory / "uvek_mapping_review.csv").open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(audit[0]))
        writer.writeheader()
        writer.writerows(audit)
    counts = dict(sorted(Counter(row["outcome"] for row in audit).items()))
    (output_directory / "uvek_mapping_review_summary.json").write_text(json.dumps(counts, indent=2) + "\n")
    statistics = {
        "source_entries": len(audit),
        "selected": len(decisions),
        "selected_previous_confidence_below_0.35": sum(row["previous_confidence"] < 0.35 for row in decisions),
        "selected_previous_confidence_below_0.5": sum(row["previous_confidence"] < 0.5 for row in decisions),
        "selected_deprecated_targets": sum(row["target"]["name"].lower().startswith("xx") for row in decisions),
        "selected_geographic_proxies": sum(row["target"]["location"] != row["source"]["location"] for row in decisions),
        "branch_methods_selected": dict(sorted(Counter(row["branch_method"] for row in decisions).items())),
    }
    (output_directory / "uvek_mapping_review_statistics.json").write_text(json.dumps(statistics, indent=2) + "\n")
    report = openpyxl.Workbook(write_only=True)
    summary = report.create_sheet("Summary")
    summary.append(["Outcome", "Count"])
    for outcome, count in counts.items():
        summary.append([outcome, count])
    for title, prefix in (
        ("Accepted", "accept_"),
        ("Retained", "retain_"),
        ("Rejected branch", "reject_"),
        ("Needs review", "review_"),
    ):
        sheet = report.create_sheet(title)
        sheet.freeze_panes = "A2"
        for column in ("A", "B", "J", "K", "L", "N"):
            sheet.column_dimensions[column].width = 55
        for column in ("C", "D", "F", "G", "H", "I"):
            sheet.column_dimensions[column].width = 14
        sheet.column_dimensions["E"].width = 32
        sheet.column_dimensions["M"].width = 24
        sheet.append(list(audit[0]))
        row_count = 1
        for row in sorted(
            audit, key=lambda row: (row["previous_confidence"], row["name"], row["reference product"], row["location"])
        ):
            if not row["outcome"].startswith(prefix):
                continue
            cells = []
            for value in row.values():
                cell = openpyxl.cell.WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cells.append(cell)
            sheet.append(cells)
            row_count += 1
        sheet.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(audit[0]))}{row_count}"
    report.save(output_directory / "uvek_mapping_review.xlsx")
    print(json.dumps(counts, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("resource", type=Path)
    parser.add_argument("output_directory", type=Path)
    parser.add_argument("--revision", required=True)
    arguments = parser.parse_args()
    review(arguments.workbook, arguments.resource, arguments.output_directory, arguments.revision)
