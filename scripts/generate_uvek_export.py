"""Generate a complete catalog disposition without approving similarity guesses."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from functools import lru_cache
from pathlib import Path

from brightpath.background.uvek_export import IDENTITY_FIELDS, RESOURCE_NAME, identity
from brightpath.migrations.schema import (
    MIGRATION_SCHEMA,
    MIGRATION_SCHEMA_VERSION,
    dumps_migration_resource,
    intern_review,
    normalize_migration_resource,
    validate_migration_resource,
)
from brightpath.units import normalize_unit

LINEAGE_METHOD = "publisher_product_correspondence"
LINEAGE_PATH = "docs/uvek_export_lineage_evidence.json"
CORRESPONDENCE_URL = (
    "https://support.ecoinvent.org/hubfs/Knowledge%20Base/Database/Releases/"
    "correspondence_file_intermediate-exchangesv2.2_to_v3.0_20130904.xlsx"
)
CORRESPONDENCE_SHA256 = "cd598b9c357a1c0f55ca5e00fc9861a753cbdce772a07dd43e8acffcb8f2b959"
CORRESPONDENCE_SHEET = "correspondance sheet_corrected"
CHANGE_COLUMNS = (
    "Geographical area changed to:",
    "Reference product and amount changed",
    "Unit changed to:",
    "Dataset deleted (replaced by v2.2 ID)",
    "Product no longer used as input",
    "New reference product added",
    "Input turned into reference product",
)
CORRESPONDENCE_COLUMNS = (
    "SINGLE-OUTPUT DATASET-ID v2.2",
    "MULTI-OUTPUT DATASET-ID v2.2",
    "PRODUCT NAME v2.2",
    "ACTIVITY NAME v3",
    "PRODUCT NAME v3",
    "UUID PRODUCT NAME v3",
    "REFERENCE PRODUCT",
    "Unit",
    "Location",
    *CHANGE_COLUMNS,
)
REVIEW_ROLE = re.compile(
    r"\b(?:waste\w*|scrap|disposal|treatment|recycl\w*|landfill|incinerat\w*|used|spent|sludge|residues?|sewage)\b",
    re.IGNORECASE,
)


def label(value):
    """Normalize case and whitespace without deleting functional qualifiers."""
    return " ".join(str(value or "").casefold().split())


def legacy_export_fingerprint(resource):
    """Bind additions to the pre-existing rules, reviews and supporting recipes."""
    rules = [rule for rule in resource["replace"] if rule.get("mapping_method") != LINEAGE_METHOD]
    reviews = {rule["review_ref"]: resource["reviews"][rule["review_ref"]] for rule in rules}
    value = {"replace": rules, "reviews": reviews, "recipes": resource["recipes"]}
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, allow_nan=False).encode()).hexdigest()


def lineage_catalogs(root):
    """Load exact catalogs and their byte digests for authoring validation."""
    directory = root / "brightpath/data/export/reference_catalogs"
    names = ("uvek__2025__cutoff.json", "uvek__2026__cutoff.json", "ecoinvent__3.12__cutoff.json")
    catalogs = {name: load(directory / name) for name in names}
    digests = {name: hashlib.sha256((directory / name).read_bytes()).hexdigest() for name in names}
    return catalogs, digests


def publisher_row_rejection(source, row, targets):
    """Reject ambiguous reference bases before considering a current target."""
    if label(source[0]) != label(row["PRODUCT NAME v2.2"]) or label(source[1]) != label(row["PRODUCT NAME v2.2"]):
        return "source_label_or_product_mismatch", None
    if source[2:] != (row["Location"], normalize_unit(row["Unit"])):
        return "source_location_or_unit_mismatch", None
    legacy_id = row["SINGLE-OUTPUT DATASET-ID v2.2"]
    if (
        isinstance(legacy_id, bool)
        or not isinstance(legacy_id, (int, float))
        or not math.isfinite(legacy_id)
        or legacy_id <= 0
        or int(legacy_id) != legacy_id
    ):
        return "missing_legacy_dataset_id", None
    if row["MULTI-OUTPUT DATASET-ID v2.2"] not in (None, ""):
        return "multioutput_requires_review", None
    if row["REFERENCE PRODUCT"] != "x":
        return "reference_product_role_requires_review", None
    if any(row[column] not in (None, "") for column in CHANGE_COLUMNS):
        return "publisher_change_flag_requires_review", None
    target_key = (label(row["ACTIVITY NAME v3"]), label(row["PRODUCT NAME v3"]), source[2], source[3])
    matches = targets.get(target_key, ())
    if len(matches) != 1:
        return "current_target_missing_or_ambiguous", None
    target = matches[0]
    if target.get("flow_id") != row["UUID PRODUCT NAME v3"] or not target.get("flow_id"):
        return "current_product_uuid_mismatch", None
    if REVIEW_ROLE.search(" ".join((source[0], target["name"], target["reference_product"]))):
        return "waste_or_secondary_product_role_requires_review", None
    return "publisher_correspondence", target


def extract_publisher_lineage(workbook, source_catalog, target_catalog, protected_sources):
    """Review exact legacy labels against single-reference publisher rows."""
    from openpyxl import load_workbook

    targets = defaultdict(list)
    for row in target_catalog["technosphere"]:
        targets[label(row["name"]), label(row["reference_product"]), row["location"], row["unit"]].append(row)
    book = load_workbook(workbook, read_only=True, data_only=True)
    rows_by_source = defaultdict(list)
    try:
        sheet = book[CORRESPONDENCE_SHEET]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        if len(headers) != len(set(headers)) or not set(CORRESPONDENCE_COLUMNS) <= set(headers):
            raise ValueError("The publisher correspondence workbook has missing or duplicated columns.")
        for number, values in enumerate(rows, 2):
            row = dict(zip(headers, values, strict=True))
            if not all(row[key] for key in ("PRODUCT NAME v2.2", "Unit", "Location")):
                continue
            key = label(row["PRODUCT NAME v2.2"]), row["Location"], normalize_unit(row["Unit"])
            rows_by_source[key].append((number, {column: row[column] for column in CORRESPONDENCE_COLUMNS}))
    finally:
        book.close()
    decisions, counts = [], Counter()
    for source_row in source_catalog["technosphere"]:
        source = source_row["name"], source_row["reference_product"], source_row["location"], source_row["unit"]
        if source in protected_sources:
            counts["existing_review_preserved"] += 1
            continue
        rows = rows_by_source[label(source[0]), source[2], source[3]]
        if not rows:
            counts["no_exact_legacy_label"] += 1
            continue
        if len(rows) != 1:
            counts["ambiguous_legacy_rows"] += 1
            continue
        number, row = rows[0]
        disposition, target = publisher_row_rejection(source, row, targets)
        counts[disposition] += 1
        if target is not None:
            decisions.append(
                {
                    "source": dict(zip(IDENTITY_FIELDS, source, strict=True)),
                    "target": dict(
                        zip(
                            IDENTITY_FIELDS,
                            (target["name"], target["reference_product"], target["location"], target["unit"]),
                            strict=True,
                        )
                    ),
                    "worksheet_row": number,
                    "publisher_row": row,
                }
            )
    decisions.sort(key=lambda row: identity(row["source"]))
    return decisions, dict(sorted(counts.items()))


def build_lineage_evidence(root, workbook):
    """Derive an auditable cohort from the pinned public publisher workbook."""
    if hashlib.sha256(workbook.read_bytes()).hexdigest() != CORRESPONDENCE_SHA256:
        raise ValueError("Publisher workbook checksum changed; review a new source revision before regenerating.")
    catalogs, digests = lineage_catalogs(root)
    resource = load(root / "brightpath/data/migrations/uvek" / RESOURCE_NAME)
    protected = {
        identity(rule["source"]) for rule in resource["replace"] if rule.get("mapping_method") != LINEAGE_METHOD
    }
    decisions, dispositions = extract_publisher_lineage(
        workbook, catalogs["uvek__2025__cutoff.json"], catalogs["ecoinvent__3.12__cutoff.json"], protected
    )
    return {
        "schema_version": 1,
        "generator": "scripts/generate_uvek_export.py",
        "method": LINEAGE_METHOD,
        "source": {
            "publisher": "ecoinvent",
            "title": "Correspondence File for Intermediate Exchanges, ecoinvent v2.2 to v3.01 (2013-09-04)",
            "url": CORRESPONDENCE_URL,
            "sha256": CORRESPONDENCE_SHA256,
            "sheet": CORRESPONDENCE_SHEET,
            "license": "Publisher terms apply; no explicit reuse license is declared in the workbook. Correspondence metadata only; no inventories copied.",
        },
        "catalog_sha256": digests,
        "legacy_export_sha256": legacy_export_fingerprint(resource),
        "dispositions": dispositions,
        "correspondences": decisions,
        "scientific_equivalence": False,
    }


def lineage_decisions(root, evidence):
    """Validate authoring evidence against current exact catalogs and legacy rules."""
    catalogs, digests = lineage_catalogs(root)
    resource = load(root / "brightpath/data/migrations/uvek" / RESOURCE_NAME)
    if evidence.get("schema_version") != 1 or evidence.get("method") != LINEAGE_METHOD:
        raise ValueError("Unsupported publisher-lineage evidence schema.")
    if (
        evidence["source"]["sha256"] != CORRESPONDENCE_SHA256
        or evidence["source"]["url"] != CORRESPONDENCE_URL
        or evidence["source"]["sheet"] != CORRESPONDENCE_SHEET
    ):
        raise ValueError("Publisher-lineage evidence does not name the pinned source workbook.")
    if evidence["catalog_sha256"] != digests or evidence["legacy_export_sha256"] != legacy_export_fingerprint(resource):
        raise ValueError("Publisher-lineage evidence is stale for the catalogs or pre-existing export rules.")
    protected = {
        identity(rule["source"]) for rule in resource["replace"] if rule.get("mapping_method") != LINEAGE_METHOD
    }
    sources = {
        (r["name"], r["reference_product"], r["location"], r["unit"])
        for r in catalogs["uvek__2025__cutoff.json"]["technosphere"]
    }
    targets = defaultdict(list)
    for row in catalogs["ecoinvent__3.12__cutoff.json"]["technosphere"]:
        targets[label(row["name"]), label(row["reference_product"]), row["location"], row["unit"]].append(row)
    decisions, row_numbers = {}, set()
    for entry in evidence["correspondences"]:
        source, number = identity(entry["source"]), entry["worksheet_row"]
        if source not in sources or source in protected or source in decisions or number in row_numbers:
            raise ValueError("Publisher-lineage evidence has an unknown, protected or duplicate source/row.")
        if isinstance(number, bool) or not isinstance(number, int) or number < 2:
            raise ValueError("Invalid publisher worksheet row.")
        disposition, target = publisher_row_rejection(source, entry["publisher_row"], targets)
        if target is None:
            raise ValueError(f"Publisher-lineage evidence is not an approved correspondence: {disposition}")
        expected = target["name"], target["reference_product"], target["location"], target["unit"]
        if identity(entry["target"]) != expected:
            raise ValueError("Publisher-lineage target disagrees with the exact current catalog identity.")
        row_numbers.add(number)
        decisions[source] = {
            "target": dict(entry["target"]),
            "factor": 1.0,
            "classification": "correspondence",
            "mapping_method": LINEAGE_METHOD,
            "rationale": "BrightPath reference-function correspondence using the publisher's legacy nomenclature table: exact UVEK label, unit and geography; reference-product row marked x without change flags; exact current activity/product/location/unit and stable product UUID. This is not a publisher-certified UVEK mapping or equality of background coefficients, allocation or LCIA results.",
            "evidence": [
                f"{CORRESPONDENCE_URL}; sheet {CORRESPONDENCE_SHEET}; row {number}",
                f"{LINEAGE_PATH}; publisher product UUID {entry['publisher_row']['UUID PRODUCT NAME v3']}",
            ],
        }
    return decisions


def load(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, indent=2, ensure_ascii=False, sort_keys=True, allow_nan=False) + "\n",
        encoding="utf-8",
        newline="\n",
    )


def clean(value):
    return " ".join(re.sub(r"\s*\{[^{}]+\}\s*[US]?\s*$", "", value).casefold().split())


def source_key(name, location, unit):
    return clean(name), location, normalize_unit(unit)


@lru_cache(maxsize=65536)
def tokens(value):
    """Index distinctive supplier terms rather than process boilerplate."""
    return frozenset(re.findall(r"[a-z0-9]+", clean(value))) - {
        "at",
        "in",
        "of",
        "for",
        "the",
        "and",
        "plant",
        "production",
        "market",
        "group",
        "from",
        "with",
        "by",
    }


def candidate_index(targets):
    index = defaultdict(set)
    for target in targets:
        for token in tokens(target[0] + " " + target[1]):
            index[target[3], token].add(target)
    return index


def ranked_candidates(source, index):
    """Suggest review candidates; ranking scores are not mapping confidence."""
    source_tokens = tokens(source[0] + " " + source[1])
    candidates = set().union(*(index[source[3], token] for token in source_tokens)) if source_tokens else set()

    def score(target):
        target_tokens = tokens(target[0] + " " + target[1])
        similarity = len(source_tokens & target_tokens) / max(1, len(source_tokens | target_tokens))
        return -similarity, -(source[2] == target[2]), target

    return sorted(candidates, key=score)[:5]


def extract_evidence(args):
    """Reduce user-supplied public PV provenance to standalone rule evidence."""
    source_path = Path(args.pv_suppliers)
    candidates = []
    for row in csv.DictReader(source_path.open()):
        candidates.append(
            {
                "source_name": row["source_supplier"],
                "source_location": row["source_location"],
                "source_unit": row["source_unit"],
                "target": {
                    "name": row["ecoinvent_activity"],
                    "reference product": row["ecoinvent_reference_product"],
                    "location": row["ecoinvent_location"],
                    "unit": row["ecoinvent_unit"],
                },
                "factor": float(row["amount_factor"]),
                "classification": "correspondence" if row["mapping_type"] == "Exact correspondence" else "proxy",
                "rationale": row["rationale"],
                "evidence": [
                    "IEA PVPS T12-33:2026 https://doi.org/10.69766/WJTE1771",
                    f"{source_path.name}; sheets: {row['source_sheets']}",
                ],
            }
        )
    aliases = []
    for entry in load(args.pv_audit):
        if entry["method"] == "explicit_catalog_label_alias" and entry["type"] == "technosphere":
            aliases.append({"original": entry["source"], "canonical": entry["target"]})
    from brightpath import (
        BackgroundContext,
        BiosphereProfile,
        BrightwayInventory,
        FormatProfile,
        InventoryContext,
        TechnosphereProfile,
    )

    context = InventoryContext(
        FormatProfile("brightway_excel"),
        BackgroundContext(TechnosphereProfile("ecoinvent", "3.12", "cutoff"), BiosphereProfile("ecoinvent", "3.12")),
    )
    datasets = BrightwayInventory.from_excel(args.pv_workbook, context=context).data
    naming = {entry["source_id"]: entry for entry in load(args.pv_naming)}
    recipes = []
    for entry in load(args.pv_summary)["externalized_supporting_datasets"]:
        label = naming[entry["source_id"]]
        matches = [
            dataset
            for dataset in datasets
            if dataset["name"] == label["name"] and dataset["location"] == label["location"]
        ]
        if len(matches) != 1:
            raise ValueError(f"Cannot identify public support recipe: {entry['source_id']}")
        dataset = matches[0]
        inputs = []
        for exchange in dataset["exchanges"]:
            if exchange["type"] == "production":
                continue
            inputs.append(
                {key: value for key, value in exchange.items() if key not in {"input", "database", "code", "output"}}
            )
        recipes.append(
            {
                "source": dict(zip(IDENTITY_FIELDS, entry["native_supplier"], strict=True)),
                "recipe": {
                    "name": dataset["name"],
                    "reference product": dataset["reference product"],
                    "location": dataset["location"],
                    "unit": dataset["unit"],
                    "inputs": inputs,
                    "rationale": dataset.get("comment") or inputs[0].get("comment", "Public supporting recipe."),
                    "evidence": [
                        dataset.get("source", "IEA PVPS T12-33:2026 supporting inventories"),
                        entry["source_id"],
                    ],
                },
            }
        )
    paths = [source_path, Path(args.pv_audit), Path(args.pv_workbook), Path(args.pv_naming), Path(args.pv_summary)]
    return {
        "schema_version": 1,
        "candidates": candidates,
        "aliases": aliases,
        "recipes": recipes,
        "input_sha256": {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in paths},
    }


def build(root, evidence, lineage=None):
    directory = root / "brightpath/data"
    source_catalog = load(directory / "export/reference_catalogs/uvek__2025__cutoff.json")
    target_catalog = load(directory / "export/reference_catalogs/ecoinvent__3.12__cutoff.json")
    sources = {
        identity({**row, "reference product": row["reference_product"]}): row for row in source_catalog["technosphere"]
    }
    targets = {
        identity({**row, "reference product": row["reference_product"]}) for row in target_catalog["technosphere"]
    }
    candidates = defaultdict(list)
    overrides = load(root / "docs/uvek_export_overrides.json")
    for candidate in [*evidence["candidates"], *overrides["candidates"]]:
        candidates[source_key(candidate["source_name"], candidate["source_location"], candidate["source_unit"])].append(
            candidate
        )
    aliases = defaultdict(set)
    for alias in evidence["aliases"]:
        original, canonical = alias["original"], alias["canonical"]
        aliases[source_key(canonical["name"], canonical["location"], canonical["unit"])].add(
            source_key(original["name"], original["location"], original["unit"])
        )
    recipes, recipe_rules = {}, {}
    for entry in evidence["recipes"]:
        source = identity(entry["source"])
        if source not in sources:
            raise ValueError(f"Recipe source is not in UVEK catalog: {source}")
        recipe_id = hashlib.sha256(json.dumps(source).encode()).hexdigest()[:16]
        recipes[recipe_id] = entry["recipe"]
        recipe_rules[source] = {
            "recipe": recipe_id,
            "classification": "proxy",
            "rationale": entry["recipe"]["rationale"],
            "evidence": entry["recipe"]["evidence"],
        }
    rules, review = [], []
    review_resource = {"reviews": {}}
    lineage = lineage_decisions(root, lineage) if lineage is not None else {}
    target_index = candidate_index(targets)
    forward = normalize_migration_resource(load(directory / "migrations/uvek/ecoinvent-to-uvek-2025.json"))
    suggestions = defaultdict(set)
    for rule in forward["replace"]:
        if identity(rule["source"]) in targets:
            suggestions[identity(rule["target"])].add(identity(rule["source"]))
    for source in sorted(sources):
        source_record = dict(zip(IDENTITY_FIELDS, source, strict=True))
        key = source_key(source[0], source[2], source[3])
        matches = list(candidates[key])
        for alias in sorted(aliases[key]):
            matches.extend(candidates[alias])
        matches = [candidate for candidate in matches if identity(candidate["target"]) in targets]
        signatures = {(identity(candidate["target"]), candidate["factor"]) for candidate in matches}
        decision = recipe_rules.get(source)
        if decision is None and len(signatures) == 1:
            selected = matches[0]
            decision = {field: selected[field] for field in ("target", "factor", "classification", "rationale")}
            decision["evidence"] = sorted({reference for candidate in matches for reference in candidate["evidence"]})
            if any(candidate["classification"] == "proxy" for candidate in matches):
                decision["classification"] = "proxy"
        if source in lineage:
            if decision is not None or matches:
                raise ValueError(f"Publisher lineage overlaps a pre-existing evidence decision: {source}")
            decision = lineage[source]
        if decision is not None:
            rule_id = "uvek-ei312-" + hashlib.sha256(json.dumps(source).encode()).hexdigest()[:16]
            rule = {
                "id": rule_id,
                "source": source_record,
                "target": {"recipe_ref": decision["recipe"]} if "recipe" in decision else decision["target"],
                "classification": decision["classification"],
                "review_ref": intern_review(review_resource, {key: decision[key] for key in ("rationale", "evidence")}),
            }
            if "factor" in decision:
                rule["conversion_factor"] = decision["factor"]
            if "mapping_method" in decision:
                rule["mapping_method"] = decision["mapping_method"]
            rules.append(rule)
            review.append(
                {
                    "source": source_record,
                    "disposition": "recipe" if "recipe" in decision else decision["classification"],
                    "rule_id": rule_id,
                }
            )
        else:
            reason = (
                "Conflicting source-backed targets or factors require adjudication."
                if len(signatures) > 1
                else "No reviewed directional correspondence; similarity and reversed forward rules are not approval."
            )
            review.append(
                {
                    "source": source_record,
                    "disposition": "unresolved",
                    "reason": reason,
                    "candidates": [
                        dict(zip(IDENTITY_FIELDS, target, strict=True))
                        for target in ranked_candidates(source, target_index)
                    ],
                    "candidate_method": "unit-constrained token ranking, not approval",
                    "reverse_candidate_count": len(suggestions[source]),
                }
            )
    for recipe in recipes.values():
        for exchange in recipe["inputs"]:
            if exchange["type"] == "technosphere" and identity(exchange) not in targets:
                raise ValueError(f"Public recipe has a non-background dependency: {identity(exchange)}")
    coverage = {
        "source_identities": len(sources),
        "mapped_identities": len(rules),
        "unresolved_identities": len(sources) - len(rules),
        "dispositions": dict(Counter(entry["disposition"] for entry in review)),
    }
    resource = {
        "schema": MIGRATION_SCHEMA,
        "schema_version": MIGRATION_SCHEMA_VERSION,
        "status": "active",
        "axis": "technosphere",
        "quality": "reviewed",
        "name": "uvek-2025-to-ecoinvent-3.12-cutoff",
        "source_id": "uvek-2025-cutoff",
        "target_id": "ecoinvent-3.12-cutoff",
        "source_profiles": [{"family": "uvek", "version": "2025", "system_model": "cutoff"}],
        "target_profile": {"family": "ecoinvent", "version": "3.12", "system_model": "cutoff"},
        "replace": rules,
        "disaggregate": [],
        "delete": [],
        "reverse_policy": "explicit_only",
        "reviews": review_resource["reviews"],
        "recipes": recipes,
        "coverage": coverage,
    }
    validate_migration_resource(resource)
    return resource, {"coverage": coverage, "review": review, "scientific_equivalence": False}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--evidence", type=Path)
    parser.add_argument("--lineage-workbook", type=Path)
    parser.add_argument("--lineage-evidence", type=Path)
    parser.add_argument("--review-output", type=Path)
    for flag in ("pv-suppliers", "pv-audit", "pv-workbook", "pv-naming", "pv-summary"):
        parser.add_argument("--" + flag, type=Path)
    args = parser.parse_args()
    evidence_path = args.evidence or args.root / "docs/uvek_export_evidence.json"
    if args.pv_suppliers:
        write(evidence_path, extract_evidence(args))
    lineage_path = args.lineage_evidence or args.root / LINEAGE_PATH
    if args.lineage_workbook:
        write(lineage_path, build_lineage_evidence(args.root, args.lineage_workbook))
    if not lineage_path.is_file():
        parser.error("Publisher-lineage evidence is missing; provide --lineage-workbook to regenerate it.")
    lineage = load(lineage_path)
    resource, review = build(args.root, load(evidence_path), lineage)
    path = args.root / "brightpath/data/migrations/uvek" / RESOURCE_NAME
    path.write_text(dumps_migration_resource(resource), encoding="utf-8", newline="\n")
    write(args.review_output or args.root / "docs/uvek_export_review.json", review)
    manifest_path = args.root / "brightpath/data/migrations/RESOURCE_MANIFEST.json"
    manifest = load(manifest_path)
    relative = "uvek/" + RESOURCE_NAME
    manifest["resources"] = [entry for entry in manifest["resources"] if entry["path"] != relative]
    manifest["resources"].append(
        {
            "path": relative,
            "name": resource["name"],
            "source_id": resource["source_id"],
            "target_id": resource["target_id"],
            "schema_version": resource["schema_version"],
            "status": "active",
            "licenses": [],
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "size": path.stat().st_size,
        }
    )
    manifest["resources"].sort(key=lambda entry: entry["path"])
    write(manifest_path, manifest)
    print(json.dumps(resource["coverage"], indent=2))


if __name__ == "__main__":
    main()
